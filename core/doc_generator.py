"""
Word文档生成模块
负责将内容转换为格式化的Word文档
"""
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from typing import Optional, Callable
from pathlib import Path
from datetime import datetime
import re


class DocGenerator:
    """Word文档生成器"""
    
    def __init__(self, config: dict, log_callback: Optional[Callable] = None):
        self.config = config
        self.log = log_callback if log_callback else print
        self.word_config = config.get('word_format', {})
        self.font_name = self.word_config.get('font', '微软雅黑')
        self.font_size = self.word_config.get('font_size', 12)
        self.line_space = self.word_config.get('line_space', '1.5倍')
        self.add_cover = self.word_config.get('add_cover', False)
        self.add_toc = self.word_config.get('add_toc', False)
    
    def create_document(self, title: str, content: str, output_path: str) -> bool:
        """
        创建Word文档
        """
        try:
            self.log("INFO", f"开始生成文档: {title}")
            doc = Document()
            
            self._setup_styles(doc)
            if self.add_cover: self._add_cover_page(doc, title)
            if self.add_toc: self._add_table_of_contents(doc)
            
            self._add_title(doc, title)
            self._add_content(doc, content)
            
            output_file = Path(output_path)
            output_file.parent.mkdir(parents=True, exist_ok=True)
            
            # **[FIXED]** 添加日志以确认绝对保存路径
            self.log("INFO", f"正在将文档保存到绝对路径: {output_file.resolve()}")
            
            doc.save(str(output_file))
            return True
            
        except Exception as e:
            import traceback
            self.log("ERROR", f"生成文档失败: {str(e)}")
            self.log("DEBUG", traceback.format_exc())
            return False

    def _setup_styles(self, doc: Document):
        try:
            style = doc.styles['Normal']
            font = style.font
            font.name = self.font_name
            font.size = Pt(self.font_size)
            style.element.rPr.rFonts.set(qn('w:eastAsia'), self.font_name)
            
            p_format = style.paragraph_format
            if self.line_space == '1.5倍': p_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            elif self.line_space == '2倍行距': p_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
            else: p_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            
        except Exception as e:
            self.log("WARNING", f"设置样式失败: {str(e)}")
    
    def _add_cover_page(self, doc: Document, title: str):
        try:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(title)
            run.font.size = Pt(26)
            run.font.bold = True
            for _ in range(10): doc.add_paragraph()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(f"生成日期: {datetime.now().strftime('%Y-%m-%d')}")
            run.font.size = Pt(14)
            doc.add_page_break()
        except Exception as e:
            self.log("WARNING", f"添加封面页失败: {e}")

    def _add_table_of_contents(self, doc: Document):
        try:
            paragraph = doc.add_paragraph()
            run = paragraph.add_run()
            fldChar = OxmlElement('w:fldChar')
            fldChar.set(qn('w:fldCharType'), 'begin')
            run._r.append(fldChar)
            instrText = OxmlElement('w:instrText')
            instrText.set(qn('xml:space'), 'preserve')
            instrText.text = 'TOC \\o "1-3" \\h \\z \\u'
            run._r.append(instrText)
            fldChar = OxmlElement('w:fldChar')
            fldChar.set(qn('w:fldCharType'), 'end')
            run._r.append(fldChar)
            doc.add_page_break()
            self.log("INFO", "已添加目录 (需在Word中手动更新)")
        except Exception as e:
            self.log("WARNING", f"添加目录失败: {e}")

    def _add_title(self, doc: Document, title: str):
        try:
            doc.add_heading(title, level=0)
        except Exception as e:
            self.log("WARNING", f"添加主标题失败: {e}")

    def _add_content(self, doc: Document, content: str):
        try:
            if not content: return
            
            paragraphs = content.split('\n')
            for para_text in paragraphs:
                if not para_text.strip(): continue

                level = self._detect_heading_level(para_text.strip())
                if level > 0:
                    doc.add_heading(para_text.strip(), level=level)
                else:
                    doc.add_paragraph(para_text)
        except Exception as e:
            self.log("ERROR", f"添加正文内容失败: {str(e)}")

    def _detect_heading_level(self, text: str) -> int:
        if re.match(r'^[一二三四五六七八九十]+[、．]', text) or text.startswith('# '): return 1
        if re.match(r'^[（(][一二三四五六七八九十\d]+[）)]', text) or re.match(r'^\d+[、．.]', text) or text.startswith('## '): return 2
        if re.match(r'^\d+[)）]', text) or text.startswith('### '): return 3
        return 0